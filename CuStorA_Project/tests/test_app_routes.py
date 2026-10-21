"""Dashboard data endpoints (run against the shipped analysis outputs)."""

import pandas as pd
import pytest

import app as custora

BUSINESS = "data/processed/customers_business.csv"


@pytest.fixture(scope="module")
def client():
    return custora.app.test_client()


def test_every_state_lists_its_customers(client):
    """Regression: state buttons used to filter only the top 20 customers, so most states looked empty."""
    states = pd.read_csv(BUSINESS)["CustomerState"].value_counts()
    assert len(states) >= 2
    for state, count in states.items():
        data = client.get("/customers/list.json", query_string={"state": state}).get_json()
        assert data["total"] == count
        assert data["shown"] == min(count, 50) and data["shown"] > 0
        assert {r["CustomerState"] for r in data["rows"]} == {state}


def test_customer_search_covers_everyone_not_just_the_top_rows(client):
    ids = pd.read_csv(BUSINESS)["CustomerID"].astype(str)
    target = ids.iloc[-1]
    data = client.get("/customers/list.json", query_string={"q": target}).get_json()
    assert target in [str(r["CustomerID"]) for r in data["rows"]]


def test_old_3d_constellation_is_gone(client):
    assert client.get("/constellation.json").status_code == 404


def test_matrix_endpoint_covers_every_customer_with_consistent_arrays(client):
    data = client.get("/api/matrix.json").get_json()
    n = len(pd.read_csv(BUSINESS))
    assert data["n"] == n and data["plotted"] == len(data["id"]) <= n
    assert len(data["p"]) == len(data["v"]) == len(data["l"]) == len(data["q"]) == len(data["id"])
    assert all(0 <= p <= 1 for p in data["p"])
    assert set(data["q"]) <= set(range(len(data["quadrants"])))
    assert 0 < data["risk_cut"] < 1 and data["value_cut"] > 0


def test_quadrant_filter_matches_the_dashboard_counts(client):
    total = 0
    for key in ("save_now", "protect", "automate", "healthy"):
        data = client.get("/customers/list.json", query_string={"quadrant": key, "limit": 1}).get_json()
        total += data["total"]
    assert total == len(pd.read_csv(BUSINESS))


def test_list_sorting_by_expected_loss_is_descending_and_pages(client):
    first = client.get("/customers/list.json", query_string={"sort": "loss", "limit": 20}).get_json()
    losses = [r["ExpectedLoss"] for r in first["rows"]]
    assert losses == sorted(losses, reverse=True)
    second = client.get("/customers/list.json", query_string={"sort": "loss", "limit": 20, "offset": 20}).get_json()
    assert {r["CustomerID"] for r in first["rows"]}.isdisjoint({r["CustomerID"] for r in second["rows"]})
    assert second["offset"] == 20


def test_call_list_export_is_a_csv_of_the_top_customers(client):
    resp = client.get("/customers/export.csv", query_string={"limit": 7, "sort": "loss"})
    assert resp.status_code == 200 and resp.mimetype == "text/csv"
    lines = resp.data.decode().strip().splitlines()
    assert len(lines) == 8 and lines[0].startswith("CustomerID,")


def test_call_list_excel_export_is_formatted_and_safe(client):
    import io

    from openpyxl import load_workbook

    resp = client.get("/customers/export.xlsx", query_string={"limit": 7, "sort": "loss"})
    assert resp.status_code == 200 and resp.mimetype.endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Call list", "About this list"]
    ws = wb["Call list"]
    assert ws.max_row == 8 and ws.freeze_panes == "C2" and ws.auto_filter.ref
    assert ws["C2"].number_format == "0.0%" and 0 <= ws["C2"].value <= 1
    losses = [ws.cell(row=r, column=5).value for r in range(2, 9)]
    assert losses == sorted(losses, reverse=True)
    # nothing in the text columns may be stored as a formula
    assert all(c.data_type != "f" for row in ws.iter_rows(min_row=2) for c in row)


def test_export_neutralises_spreadsheet_formulas(client):
    assert custora._csv_safe("=HYPERLINK(\"http://x\")").startswith("'")
    assert custora._csv_safe("+1") == "'+1" and custora._csv_safe("safe") == "safe"


def test_dashboard_and_evidence_pages_render(client):
    for url in ("/", "/evidence", "/upload"):
        assert client.get(url).status_code == 200


def test_sidebar_follows_the_theme_instead_of_being_forced_dark():
    """Regression: theme.css used to force a dark sidebar gradient in the light theme."""
    root = custora.app.root_path
    theme = open(f"{root}/static/css/theme.css").read()
    app_css = open(f"{root}/static/css/app.css").read()
    assert "--sidebar-bg" not in theme
    light, dark = app_css.split('[data-theme="dark"]', 1)
    assert "--sidebar-bg:#FFFFFF" in light  # light theme -> light sidebar
    assert "--sidebar-bg:#0A1512" in dark.split("}")[0]  # dark theme -> dark sidebar
    for page in ("dashboard.html", "evidence.html", "upload.html", "customer/profile.html"):
        html = open(f"{root}/templates/{page}").read()
        assert "_sidebar.html" in html or "_sidebar" in html, page
